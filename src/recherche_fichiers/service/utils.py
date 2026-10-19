"""
service/utils.py

Fonctions utilitaires pour le service d'indexation et de recherche :
    - Contrôle des fichiers
    - Lecture des fichiers multi-formats (TXT, MD, PDF, DOCX, CSV)
    - Découpage (chunking) du texte
    - Formatage des résultats de recherche
"""

# --------------- #
# --- Modules --- #
# --------------- #

import re
from pathlib import Path

# Utilisation du module de logging partagé du projet
import logging
logger = logging.getLogger(__name__)

# ------------------ #
# --- Constantes --- #
# ------------------ #

CHUNK_SIZE    = 500
CHUNK_OVERLAP = 50

# ----------------- #
# --- Fonctions --- #
# ----------------- #

def list_files_by_format(
    directory: Path, 
    formats  : set[str]
) -> list[Path]:
    """
    Liste (de manière récursive) les fichiers d'un répertoire dont l'extension (ex: ".txt", ".md")  correspond à l'un des 
    formats spécifiés.

    Args:
        directory: répertoire racine dans lequel chercher les fichiers.
        formats  : ensemble d'extensions acceptées, avec le point (ex: {".txt", ".md", ".pdf", ".docx", ".csv"}).

    Returns:
        Liste triée des chemins de fichiers correspondants.
    """
    return sorted(
        p for p in directory.rglob("*")
        if  p.is_file() 
        and p.suffix.lower() in formats
    )


def read_text_file(file_path: Path) -> str:
    """
    Lit le contenu d'un fichier texte (.txt, .md) en UTF-8, en ignorant les octets invalides plutôt que de planter.
    """
    return file_path.read_text(encoding="utf-8", errors="ignore")


def read_pdf_file(file_path: Path) -> str:
    """Extrait le texte d'un fichier PDF avec pdfplumber."""
    try:
        import pdfplumber
    except ImportError as e:
        logger.error("Dépendance manquante pour PDF : %s", e)
        raise

    text_content = []
    try:
        with pdfplumber.open(file_path) as pdf:
            for i, page in enumerate(pdf.pages):
                try:
                    page_text = page.extract_text()
                    if page_text:
                        text_content.append(page_text)
                except Exception as err:
                    logger.warning("Erreur lors de la lecture de la page %d de %s: %s", i, file_path, err)
        return "\n".join(text_content)
    except Exception as e:
        logger.error("Erreur critique lors de la lecture du PDF %s: %s", file_path, e)
        return ""


def read_docx_file(file_path: Path) -> str:
    """
    Extrait le texte d'un fichier Word (.docx) avec python-docx
    (paragraphes et tableaux inclus).
    """
    try:
        import docx
    except ImportError as e:
        logger.error("Dépendance manquante pour DOCX : %s", e)
        raise

    try:
        doc = docx.Document(file_path)
        text_lines = []

        # 1. Extraction des paragraphes classiques
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                text_lines.append(paragraph.text.strip())

        # 2. Extraction du texte contenu dans les tableaux
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    text_lines.append(row_text)

        return "\n".join(text_lines)

    except Exception as e:
        logger.error("Erreur lors de la lecture du fichier DOCX %s: %s", file_path, e)
        return ""


def read_csv_file(file_path: Path) -> str:
    """
    Extrait le texte d'un fichier CSV avec pandas (conforme au Sprint 2).
    """
    try:
        import pandas as pd
    except ImportError as e:
        logger.error("Dépendance manquante pour CSV : %s", e)
        raise

    try:
        # Lecture avec pandas
        df = pd.read_csv(file_path, encoding="utf-8", on_bad_lines="skip")
        return df.to_string(index=False)
    except Exception as e:
        logger.error("Erreur lors de la lecture du fichier CSV %s: %s", file_path, e)
        return ""


def load_file_content(file_path: Path) -> str:
    """
    Charge le contenu textuel d'un fichier en choisissant le parser approprié selon son extension.
    Gère : .txt, .md, .pdf, .docx, .csv
    """
    suffix = file_path.suffix.lower()
    
    if suffix in {".txt", ".md"}:
        return read_text_file(file_path)
    elif suffix == ".pdf":
        return read_pdf_file(file_path)
    elif suffix == ".docx":
        return read_docx_file(file_path)
    elif suffix == ".csv":
        return read_csv_file(file_path)
    else:
        raise ValueError(f"Format de fichier non pris en charge : {suffix}")


def make_document_id(file_path: Path) -> str:
    """
    Génère un `document_id` stable et sans caractères problématiques (utilisable comme préfixe d'identifiant ChromaDB) 
    à partir du chemin absolu du fichier.
    """
    return re.sub(
        r"[^a-zA-Z0-9_-]", 
        "_", 
        str(file_path.resolve())
    )


def chunk_text(
    text      : str, 
    chunk_size: int = CHUNK_SIZE, 
    overlap   : int = CHUNK_OVERLAP
) -> list[str]:
    """
    Découpe un texte en chunks de `chunk_size` caractères maximum, avec un chevauchement `overlap` entre chunks successifs. 
    Le point de coupe est recherché sur un saut de ligne ou une fin de phrase proche de la limite afin d'éviter de couper 
    un mot en deux.

    Args:
        text      : texte source à découper.
        chunk_size: taille maximale d'un chunk, en caractères.
        overlap   : chevauchement entre deux chunks consécutifs, en caractères.

    Returns:
        Liste ordonnée des chunks (chaînes non vides).
    """
    text = text.strip()
    if not text:
        return []
    if len(text) <= chunk_size:
        return [text]

    chunks: list[str] = []
    start             = 0
    text_length       = len(text)

    while start < text_length:
        end = min(start + chunk_size, text_length)

        if end < text_length:
            boundary = text.rfind("\n", start, end)
            if boundary <= start:
                boundary = text.rfind(". ", start, end)
            if boundary > start:
                end = boundary + 1

        chunk = text[start:end].strip()
        if chunk:
            chunks.append(chunk)

        next_start = end - overlap
        start      = next_start if next_start > start else end

    return chunks


def format_search_results(raw_results: dict) -> list[dict]:
    """
    Convertit le résultat brut retourné par `VectorStore.search()` ou `LexicalStore.search()` en une liste de dicts 
    directement exploitables par l'UI [{content, source_path, chunk_index, score}, ...].
    """
    documents = (raw_results.get("documents") or [[]])[0]
    metadatas = (raw_results.get("metadatas") or [[]])[0]
    scores    = (raw_results.get("scores") or raw_results.get("distances") or [[]])[0]
    results   = []

    for content, metadata, score in zip(documents, metadatas, scores):
        metadata = metadata or {}
        results.append({
            "content"    : content,
            "source_path": metadata.get("source_path"),
            "chunk_index": metadata.get("chunk_index"),
            "score"      : score,
        })
    return results